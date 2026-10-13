#!/usr/bin/env python
"""Assemble the thesis into the OFFICIAL SFSU Word template.

    python scripts/build_thesis_docx.py

WHY THIS EXISTS, AND WHY IT IS NOT build_thesis.py

`build_thesis.py` makes a handsome PDF through Chrome. SF State does not accept
that. The Division of Graduate Studies requires **Microsoft Word** unless the
department approves otherwise, runs a two-stage format check, and publishes a
checklist this script was written against:

    https://grad.sfsu.edu/content/thesis-dissertation-guidelines

Rather than imitate the template's rules, this script *opens the template* and
writes into it. Margins, the four section breaks, and the page-number scheme
(no numbers on i-iv, lower-case roman from page v, arabic from body page 1) are
inherited rather than reproduced, because reproducing them is how you get a
document that looks right and fails the check.

WHAT IS INHERITED FROM THE TEMPLATE
  * 1-inch margins on all four sides
  * the section structure and its page numbering
  * the named styles: Body, Heading 2/3/4/5, Figure Title, Table Title,
    Caption, Appendix Title
  * the Table of Contents / List of Tables / List of Figures FIELDS

WHAT THIS SCRIPT STILL CANNOT DO, AND WHAT YOU MUST DO IN WORD
  1. **Update the ToC, List of Tables and List of Figures.** They are Word
     fields, and page numbers do not exist until Word has laid the document
     out. `scripts/finalize_docx.py` drives Word to do it (the Ctrl+A, F9
     step) and then checks the result; run it after every rebuild.
  2. **Run the Accessibility Checker** (Review > Check Accessibility) and clear
     every error. Alt text is set on all seven images by this script; the
     checker also wants table header rows marked, which it is set to do, but
     it is the checker's verdict that counts, not this docstring's.
  3. **Table titles come from the source.** The checklist requires a numbered
     title above every table. Each table in `writeup/*.md` is preceded by
     `<!-- table: Its title -->`; the PDF drops it as a comment and this script
     makes it the Word Table Title. A table WITHOUT one falls back to the
     nearest heading and is listed in `writeup/TABLE_TITLES_TO_FIX.md`,
     because a generated title that looks finished is worse than an obvious
     stub. The file is deleted when every table has a real title.
  4. **Fill in the Certification of Approval names** for the committee.

THE TITLE
The checklist requires the title to match the one on the approved proposal, or
a Title Change Form to be sent to culminatingexperience@sfsu.edu. The proposal
(`FaultLines_Proposal_v2.docx`) says:

    Detecting and Mapping Silent Consensus:
    Correlated Blind Spots in LLM Ensembles

which is 11 words, inside the 12-word limit. That is what this script uses.
"FaultLines" is carried as the system's name in the abstract and the body, not
as the title, so that no form is needed.
"""
from __future__ import annotations

import argparse
import re
import shutil
import sys
import urllib.request
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
WRITEUP = ROOT / "writeup"
FIGDIR = WRITEUP / "figures"
TEMPLATE = WRITEUP / "_sfsu_template.docx"
TEMPLATE_URL = ("https://grad.sfsu.edu/sites/default/files/documents/"
                "Masters-Thesis-and-Creative-Work-Template.docx")

TITLE = ("Detecting and Mapping Silent Consensus: "
         "Correlated Blind Spots in LLM Ensembles")
AUTHOR = "Shivani Bokka"
DEGREE = "Master of Science in Data Science and Artificial Intelligence"
YEAR = "2026"

MISSING_FIGURES: list = []

FRONT = "00-front-matter.md"
CHAPTERS = [
    ("ch1-introduction.md", "Introduction"),
    ("ch2-related-work.md", "Related Work"),
    ("ch3-design.md", "Design"),
    ("ch4-method.md", "Method"),
    ("ch5-results.md", "Results"),
    ("ch6-evaluation.md", "Evaluation"),
    ("ch7-limitations.md", "Limitations"),
    ("ch8-conclusion.md", "Conclusion"),
    ("ch9-references.md", "References"),
    # supporting tables moved out of the chapters on 2026-09-27 (75 -> ~33
    # in the text); SFSU places appendices after the references
    ("ch10-appendix.md", "Appendix"),
]

# markdown heading level -> template style. The template's ToC starts at
# Heading 2, so a chapter title is Heading 2 and not Heading 1.
HSTYLE = {1: "Heading 2", 2: "Heading 3", 3: "Heading 4",
          4: "Heading 5", 5: "Heading 5", 6: "Heading 5"}

INLINE = re.compile(
    r"(\*\*\*.+?\*\*\*|\*\*.+?\*\*|(?<!\*)\*(?!\s)[^*]+?\*|`[^`]+`|"
    r"\[[^\]]+\]\([^)]+\))", re.S)
IMG = re.compile(r"^!\[(?P<alt>[^\]]*)\]\((?P<src>[^)]+)\)\s*$")
CAPTION = re.compile(r"^\*\*(?:Figure|Table)\s+[0-9.]+\s*[—\-]\s*"
                     r"(?P<title>[\s\S]+?)\*\*(?P<rest>[\s\S]*)$")


# --------------------------------------------------------------------------
# inline runs
# --------------------------------------------------------------------------
def add_inline(par, text: str, bold=None, italic=None, code=False) -> None:
    """Split markdown emphasis into Word runs, RECURSIVELY.

    The first version did not recurse, so `**`n_clusters` is ordered**` matched
    the bold rule, was emitted whole, and printed its own backticks onto the
    page. Nested emphasis is common in these chapters, so each branch re-enters
    with the inner text and the accumulated formatting.

    bold/italic default to None, meaning INHERIT from the paragraph style. They
    defaulted to False, which wrote an explicit "not bold" onto every run and
    silently un-bolded every heading in the thesis: the template's Heading 2-4
    are bold, and the document rendered them as body text."""
    for piece in INLINE.split(text):
        if not piece:
            continue
        if piece.startswith("***") and piece.endswith("***"):
            add_inline(par, piece[3:-3], True, True, code)
        elif piece.startswith("**") and piece.endswith("**"):
            add_inline(par, piece[2:-2], True, italic, code)
        elif piece.startswith("`") and piece.endswith("`"):
            r = par.add_run(piece[1:-1])
            r.bold, r.italic = bold, italic
            r.font.name = "Consolas"
        elif piece.startswith("*") and piece.endswith("*") and len(piece) > 2:
            add_inline(par, piece[1:-1], bold, True, code)
        elif piece.startswith("[") and "](" in piece:
            r = par.add_run(piece[1:piece.index("](")])
            r.bold, r.italic = bold, italic
        else:
            r = par.add_run(piece)
            r.bold, r.italic = bold, italic
            if code:
                r.font.name = "Consolas"


# --------------------------------------------------------------------------
# block parsing
# --------------------------------------------------------------------------
TTITLE = re.compile(r"<!--\s*table:\s*([\s\S]+?)\s*-->")
TT_SENTINEL = "\x00TTITLE\x00"


def blocks(md: str):
    """Yield ('kind', payload) for every block. Kinds: h, p, table, code,
    quote, ul, ol, img, hr, ttitle."""
    # A table's title is written in the source as <!-- table: ... --> so the
    # PDF (python-markdown drops comments) never shows it. It becomes its own
    # block here, for the Emitter to put above the table that follows it.
    md = TTITLE.sub(lambda m: "\n" + TT_SENTINEL + " ".join(m.group(1).split())
                    + "\n", md)
    # python-markdown drops HTML comments; this parser did not, so a note left
    # in the source rendered as literal "<!-- ... -->" on the abstract page.
    md = re.sub(r"<!--[\s\S]*?-->", "", md)
    lines = md.replace("\r\n", "\n").split("\n")
    i, n = 0, len(lines)
    while i < n:
        ln = lines[i]
        s = ln.strip()
        if not s:
            i += 1; continue
        if ln.startswith(TT_SENTINEL):
            yield ("ttitle", ln[len(TT_SENTINEL):].strip()); i += 1; continue
        if s.startswith("```"):
            buf = []; i += 1
            while i < n and not lines[i].strip().startswith("```"):
                buf.append(lines[i]); i += 1
            i += 1
            yield ("code", "\n".join(buf)); continue
        if s.startswith("#"):
            lvl = len(s) - len(s.lstrip("#"))
            yield ("h", (lvl, s.lstrip("#").strip())); i += 1; continue
        if s in ("---", "***", "___"):
            yield ("hr", None); i += 1; continue
        m = IMG.match(s)
        if m:
            yield ("img", (m.group("alt"), m.group("src"))); i += 1; continue
        if s.startswith("|"):
            rows = []
            while i < n and lines[i].strip().startswith("|"):
                rows.append(lines[i].strip()); i += 1
            yield ("table", rows); continue
        if s.startswith(">"):
            buf = []
            while i < n and (lines[i].strip().startswith(">")
                             or (buf and lines[i].strip())):
                buf.append(re.sub(r"^\s*>\s?", "", lines[i])); i += 1
            yield ("quote", " ".join(x.strip() for x in buf if x.strip()))
            continue
        if re.match(r"^([-*+]|\d+\.)\s+", s):
            items, ordered = [], bool(re.match(r"^\d+\.", s))
            while i < n and re.match(r"^\s*([-*+]|\d+\.)\s+", lines[i]):
                item = re.sub(r"^\s*([-*+]|\d+\.)\s+", "", lines[i]); i += 1
                while i < n and lines[i].strip() and \
                        not re.match(r"^\s*([-*+]|\d+\.)\s+", lines[i]) and \
                        not lines[i].strip().startswith(("|", "#", ">", "```")):
                    item += " " + lines[i].strip(); i += 1
                items.append(item)
            yield ("ol" if ordered else "ul", items); continue
        buf = []
        while i < n and lines[i].strip() and \
                not lines[i].strip().startswith(("|", "#", ">", "```", "![")) \
                and lines[i].strip() not in ("---", "***", "___") and \
                not re.match(r"^\s*([-*+]|\d+\.)\s+", lines[i]):
            buf.append(lines[i].strip()); i += 1
        if buf:
            yield ("p", " ".join(buf))
        else:
            i += 1


def set_alt_text(doc, alt: str) -> None:
    """python-docx has no alt-text API. The Accessibility Checker requires it,
    so the docPr attributes are set directly on the last inline shape."""
    shp = doc.inline_shapes[-1]
    docPr = shp._inline.docPr
    docPr.set("descr", alt)
    docPr.set("title", alt[:80])


def mark_header_row(table) -> None:
    """Repeat the header row across pages AND mark it as a header, which is
    what 'Table headings have been identified' on the checklist means."""
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    tr = table.rows[0]._tr
    trPr = tr.get_or_add_trPr()
    for tag in ("w:tblHeader",):
        el = OxmlElement(tag)
        el.set(qn("w:val"), "true")
        trPr.append(el)


# --------------------------------------------------------------------------
# the template's own furniture: content controls and comments
# --------------------------------------------------------------------------
# The template carries 25 Word CONTENT CONTROLS (w:sdt) holding placeholder
# text, and 63 comment anchors. Neither is visible through python-docx's
# paragraph.runs, because an sdt's runs live in w:sdt/w:sdtContent and are not
# direct children of w:p. The first version of this script cleared runs and
# wrote new text, which left every placeholder on the page with the real text
# appended after it -- "Choose an item.Detecting and Mapping Silent Consensus".
#
# So the controls are FILLED IN PLACE rather than deleted, which keeps the
# template's own centring and spacing, and the sentences around them are left
# completely alone.
# Matched EXACTLY on the control's placeholder text, not by substring. The
# substring version had "Choose degree" (the professor's) matching "Click to
# choose Degree" (the author's), so every committee member was awarded a
# Master of Science.
CONTROLS = {
    "title of culminating experience": "TITLE",
    "choose an item.": "thesis",
    "click to choose degree": "Master of Science",
    "click to choose major: concentration":
        "Data Science and Artificial Intelligence",
    "first middle last name": "AUTHOR",
    "click to choose semester": "Fall",
    "20xx": "YEAR",
    "(example) master of science in biology: conservation biology":
        "Master of Science in Data Science and Artificial Intelligence",
    "name of the professor,": "COMMITTEE_NAME",
    "choose degree": "COMMITTEE_DEGREE",
    "click to choose the faculty rank": "COMMITTEE_RANK",
}

# the committee, in the template's order: chair first. The template has four
# signature blocks; blocks beyond this list are DELETED, not left as stubs.
# No "Dr." -- the degree follows the name ("Name, Ph.D."), and both at once
# is the double title SFSU's own sample page avoids. A bracketed value is one
# still to be confirmed, and is left visibly bracketed on purpose.
COMMITTEE = [
    # ranks from their SFSU pages (cs.sfsu.edu, faculty.sfsu.edu, 2026-09-27);
    # the 2026-27 bulletin still lists Humayoun as Assistant, which lags
    ("Shahrukh Humayoun,", "Ph.D.", "Associate Professor"),
    ("Jenna Wong,", "Ph.D.", "Associate Professor"),
]


def sdt_text(sdt) -> str:
    from docx.oxml.ns import qn
    return "".join(n.text or "" for n in sdt.iter(qn("w:t")))


def replace_sdt(sdt, text: str) -> None:
    """Swap a content control for a plain run that inherits its formatting."""
    import copy
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    first_run = next(iter(sdt.iter(qn("w:r"))), None)
    new_r = OxmlElement("w:r")
    if first_run is not None:
        rPr = first_run.find(qn("w:rPr"))
        if rPr is not None:
            new_r.append(copy.deepcopy(rPr))
    # the checklist requires all primary text to be black; the placeholders
    # are grey, and the inherited rPr brings the grey with it
    rPr = new_r.find(qn("w:rPr"))
    if rPr is None:
        rPr = OxmlElement("w:rPr")
        new_r.append(rPr)
    for old in rPr.findall(qn("w:color")):
        rPr.remove(old)
    col = OxmlElement("w:color")
    col.set(qn("w:val"), "000000")
    rPr.append(col)
    tnode = OxmlElement("w:t")
    tnode.set(qn("xml:space"), "preserve")
    tnode.text = text
    new_r.append(tnode)
    parent = sdt.getparent()
    parent.insert(list(parent).index(sdt), new_r)
    parent.remove(sdt)


def fill_content_controls(doc) -> int:
    """Fill every placeholder. Repeated controls (the author's name appears on
    the title page and again on the copyright page; the professor block repeats
    four times) are filled by order of appearance."""
    from docx.oxml.ns import qn
    n = 0
    prof_i = deg_i = rank_i = 0
    unused = []                           # paragraphs of signature blocks past the committee

    def para_of(el):
        while el is not None and el.tag != qn("w:p"):
            el = el.getparent()
        return el

    for sdt in list(doc.element.body.iter(qn("w:sdt"))):
        if any(True for _ in sdt.iter(qn("w:instrText"))):
            continue                      # a field (ToC) -- leave it alone
        key = CONTROLS.get(sdt_text(sdt).strip().lower())
        if key is None:
            continue
        if key == "TITLE":
            value = TITLE
        elif key == "AUTHOR":
            value = AUTHOR
        elif key == "YEAR":
            value = YEAR
        elif key in ("COMMITTEE_NAME", "COMMITTEE_DEGREE", "COMMITTEE_RANK"):
            col = {"COMMITTEE_NAME": 0, "COMMITTEE_DEGREE": 1, "COMMITTEE_RANK": 2}[key]
            i = (prof_i, deg_i, rank_i)[col]
            prof_i, deg_i, rank_i = (prof_i + (col == 0), deg_i + (col == 1),
                                     rank_i + (col == 2))
            if i >= len(COMMITTEE):
                p = para_of(sdt)
                if p is not None and p not in unused:
                    unused.append(p)
                continue
            value = COMMITTEE[i][col]
        else:
            value = key
        replace_sdt(sdt, value)
        n += 1
    # Drop the unused blocks and the blank spacer paragraph above each name.
    # Never a paragraph that carries a section break.
    for p in unused:
        prev = p.getprevious()
        for el in (prev, p):
            if el is None or el.getparent() is None or el.find(".//" + qn("w:sectPr")) is not None:
                continue
            if el is prev and "".join(t.text or "" for t in el.iter(qn("w:t"))).strip():
                continue                  # the spacer must actually be blank
            if any(b.get(qn("w:type")) == "page" for b in el.iter(qn("w:br"))):
                # the template's page break before the Abstract lives in the
                # LAST rank paragraph: keep the paragraph, drop only its text
                for r in list(el.iter(qn("w:r"))):
                    if not any(b.get(qn("w:type")) == "page" for b in r.iter(qn("w:br"))):
                        r.getparent().remove(r)
                for t in el.iter(qn("w:t")):
                    t.text = ""
                continue
            el.getparent().remove(el)
    return n


def strip_comments(doc) -> int:
    """Delete every comment. The template's first comment says 'Don't forget to
    delete all comments before submitting your final document', and there are
    63 anchors behind it."""
    from docx.oxml.ns import qn
    body = doc.element.body
    n = 0
    for tag in ("w:commentRangeStart", "w:commentRangeEnd"):
        for el in list(body.iter(qn(tag))):
            el.getparent().remove(el)
            n += 1
    for ref in list(body.iter(qn("w:commentReference"))):
        run = ref.getparent()
        (run.getparent() if run.tag == qn("w:r") else run).remove(
            run if run.tag == qn("w:r") else ref)
        n += 1
    for rid, rel in list(doc.part.rels.items()):
        if "comment" in rel.reltype.lower():
            del doc.part.rels[rid]
    return n


# --------------------------------------------------------------------------
# rendering
# --------------------------------------------------------------------------
class Emitter:
    def __init__(self, doc, report):
        self.doc = doc
        self.fig = 0
        self.tab = 0
        self.heading = "Introduction"
        self.report = report
        self.pending_fig = None   # (alt, src) waiting for its caption
        self.pending_ttitle = None  # from <!-- table: -->, for the next table

    def para(self, text, style="Body"):
        p = self.doc.add_paragraph(style=style)
        add_inline(p, text)
        return p

    def flush_pending_figure(self, title=None, caption=None):
        """The template wants 'Figure N. Title' ABOVE the image and the
        explanatory caption below it. Markdown puts everything after."""
        if not self.pending_fig:
            return
        alt, src = self.pending_fig
        self.pending_fig = None
        self.fig += 1
        from docx.shared import Inches
        t = self.doc.add_paragraph(style="Figure Title")
        add_inline(t, f"Figure {self.fig}. {title or alt}")
        # Without this the title stays where it is and the image moves to the
        # next page: Figure 1's title sat alone at the foot of page 35 with
        # half a page of white under it. Title, image and caption travel
        # together.
        t.paragraph_format.keep_with_next = True
        holder = self.doc.add_paragraph()
        holder.alignment = 1
        holder.paragraph_format.keep_with_next = True
        run = holder.add_run()
        path = WRITEUP / src
        if path.exists():
            # A 6-inch-wide screenshot of a tall panel is 8.6 inches high, and
            # with a title above and a caption below it no longer fits the 9
            # inches of text a 1-inch-margin page has. Word then pushed the
            # whole block to the next page and left a blank one behind, which
            # the checklist forbids. Width is therefore chosen from the aspect
            # ratio so the image is never taller than MAX_FIG_HEIGHT.
            MAX_W, MAX_H = 6.0, 7.0
            width = MAX_W
            try:
                from PIL import Image
                with Image.open(path) as im:
                    w_px, h_px = im.size
                if w_px and h_px:
                    width = min(MAX_W, MAX_H * w_px / h_px)
            except Exception:                        # noqa: BLE001
                pass
            run.add_picture(str(path), width=Inches(width))
            set_alt_text(self.doc, alt or f"Figure {self.fig}")
        else:
            run.add_text(f"[MISSING IMAGE: {src}]")
            MISSING_FIGURES.append(src)
        if caption:
            c = self.doc.add_paragraph(style="Caption")
            add_inline(c, caption)

    def emit(self, kind, payload):
        d = self.doc
        if kind == "h":
            self.flush_pending_figure()
            lvl, text = payload
            self.heading = text
            p = d.add_paragraph(style=HSTYLE.get(lvl, "Heading 5"))
            add_inline(p, text)
            # only the template's Heading 3 keeps with next; a Heading 4 was
            # left alone at the foot of page 72 with its section overleaf
            p.paragraph_format.keep_with_next = True
            # every chapter (and the references and appendix) opens a page.
            # Not the first: it follows the front matter's section break,
            # which already starts a page, and a second break would leave
            # a blank one.
            if lvl == 1:
                self.chapters = getattr(self, "chapters", 0) + 1
                if self.chapters > 1:
                    p.paragraph_format.page_break_before = True
        elif kind == "img":
            self.flush_pending_figure()
            self.pending_fig = payload
        elif kind == "p":
            if self.pending_fig:
                m = CAPTION.match(payload)
                if m:
                    title = " ".join(m.group("title").split())
                    rest = " ".join(m.group("rest").split()).strip()
                    self.flush_pending_figure(title, rest or None)
                    return
                self.flush_pending_figure()
            self.para(payload)
        elif kind == "quote":
            self.flush_pending_figure()
            p = self.para(payload)
            p.paragraph_format.left_indent = __import__(
                "docx").shared.Inches(0.5)
            for r in p.runs:
                r.italic = True
        elif kind == "code":
            self.flush_pending_figure()
            for line in payload.split("\n"):
                p = d.add_paragraph(style="Body")
                r = p.add_run(line)
                r.font.name = "Consolas"
                p.paragraph_format.space_after = 0
        elif kind in ("ul", "ol"):
            self.flush_pending_figure()
            style = "List Number" if kind == "ol" else "List Bullet"
            for item in payload:
                try:
                    p = d.add_paragraph(style=style)
                except KeyError:
                    p = d.add_paragraph(style="List Paragraph")
                add_inline(p, item)
        elif kind == "ttitle":
            self.flush_pending_figure()
            if self.pending_ttitle:
                raise ValueError("table title with no table after it: "
                                 f"{self.pending_ttitle!r}")
            self.pending_ttitle = payload
        elif kind == "table":
            self.flush_pending_figure()
            self.table(payload)
        elif kind == "hr":
            pass

    def table(self, rows):
        cells = [[c.strip() for c in r.strip().strip("|").split("|")]
                 for r in rows]
        cells = [c for c in cells
                 if not all(re.fullmatch(r":?-{2,}:?", x or "-") for x in c)]
        if not cells:
            return
        width = max(len(r) for r in cells)
        cells = [r + [""] * (width - len(r)) for r in cells]
        self.tab += 1
        if self.pending_ttitle:
            title = f"Table {self.tab}. {self.pending_ttitle}"
            self.pending_ttitle = None
        else:
            # "3.1b The data abstraction" -> "The data abstraction". The
            # section number belongs to the heading, not the table's title.
            head = re.sub(r"^[0-9]+(\.[0-9]+[a-z]?)?\s+", "",
                          self.heading).strip()
            title = f"Table {self.tab}. {head or self.heading}"
            self.report.append((self.tab, head or self.heading))
        tp = self.doc.add_paragraph(style="Table Title")
        add_inline(tp, title)
        tp.paragraph_format.keep_with_next = True
        t = self.doc.add_table(rows=len(cells), cols=width)
        t.style = "Table Grid"
        # a short table is kept on one page: every row but the last keeps with
        # the next. Long ones may break, and the marked header row repeats.
        # "Short" is by estimated HEIGHT, not row count: an 8-row table of
        # paragraph-length cells counted as short, could not fit beside its
        # title, and Word stranded the title alone on a blank page (4 pages).
        per_line = max(12, 88 // max(1, width))            # chars per cell line
        est_lines = sum(max(-(-len(re.sub(r"[*`]", "", v)) // per_line) for v in row)
                        for row in cells)
        whole = len(cells) <= 12 and est_lines <= 22
        for ri, row in enumerate(cells):
            for ci, val in enumerate(row):
                cell = t.cell(ri, ci)
                cell.text = ""
                p = cell.paragraphs[0]
                add_inline(p, val, bold=True if ri == 0 else None)
                if (whole and ri < len(cells) - 1) or ri == 0:
                    # ri == 0: the header row always stays with a data row,
                    # so the title never ends a page with no table under it
                    p.paragraph_format.keep_with_next = True
        mark_header_row(t)
        self.doc.add_paragraph(style="Body")


# --------------------------------------------------------------------------
INSTRUCTION = re.compile(
    r"Insert a (section|page) break|please delete this page|"
    r"If your work (does not )?include|Including a Preface|"
    r"Begin your abstract here|Normally the Acknowledgements|"
    r"Microsoft.s support page", re.I)


def has_sectpr(par) -> bool:
    """A paragraph can carry the section break that switches page numbering
    from roman to arabic. Deleting one of those silently renumbers the whole
    document, so they are cleared rather than removed."""
    from docx.oxml.ns import qn
    pPr = par._p.find(qn("w:pPr"))
    return pPr is not None and pPr.find(qn("w:sectPr")) is not None


def drop(par) -> None:
    par._p.getparent().remove(par._p)


def strip_template_furniture(doc) -> tuple:
    """Remove the template's own instructions, the optional Preface, and the
    List of Appendices, which this thesis does not have.

    The checklist asks for no excessive white space and no leftover
    placeholders; more to the point, a submitted thesis containing the words
    'Begin your abstract here' is the kind of thing a format check bounces."""
    removed = kept = 0

    # the optional Preface page, and the unused List of Appendices
    for title in ("Preface", "List of Appendices"):
        seen = False
        for par in list(doc.paragraphs):
            if par.style.name == "Heading 2" and par.text.strip() == title:
                seen = True
            elif seen:
                if has_sectpr(par):
                    break
                if par.style.name == "Heading 2" and par.text.strip():
                    break
            else:
                continue
            if has_sectpr(par):
                break
            drop(par); removed += 1

    # Whatever instruction paragraphs remain -- walked over the XML, not over
    # doc.paragraphs, because the template's Abstract is a BLOCK-LEVEL content
    # control and python-docx does not descend into one. A leftover pointing at
    # "Microsoft's support page" was sitting inside it, invisible to the
    # paragraph-level sweep and perfectly visible on page 6 of the PDF.
    from docx.oxml.ns import qn
    for par_el in list(doc.element.body.iter(qn("w:p"))):
        text = "".join(x.text or "" for x in par_el.iter(qn("w:t")))
        if not text.strip() or not INSTRUCTION.search(text):
            continue
        pPr = par_el.find(qn("w:pPr"))
        if pPr is not None and pPr.find(qn("w:sectPr")) is not None:
            # Everything except the properties, because the instruction text
            # ends in a HYPERLINK ("Microsoft's support page") and a hyperlink
            # is a w:hyperlink element, not a w:r. Removing only runs left the
            # link behind, alone on page 6, looking like a citation.
            for child in list(par_el):
                if child.tag != qn("w:pPr"):
                    par_el.remove(child)
            kept += 1
        else:
            par_el.getparent().remove(par_el)
            removed += 1
    return removed, kept


def collapse_blank_before_section_breaks(doc) -> int:
    """Delete the empty paragraphs that sit between the last real text and a
    section break.

    The template padded its abstract page with blank paragraphs and a trailing
    instruction. Once the instruction was removed the padding remained, which
    pushed the section-break paragraph onto a page of its own -- a genuinely
    blank page 5, sitting between the abstract and the acknowledgements, and
    the checklist says there are to be none.

    The paragraph carrying the break is kept; only the padding in front of it
    goes."""
    from docx.oxml.ns import qn
    removed = 0
    for par in list(doc.element.body.iter(qn("w:p"))):
        pPr = par.find(qn("w:pPr"))
        if pPr is None or pPr.find(qn("w:sectPr")) is None:
            continue
        prev = par.getprevious()
        while prev is not None and prev.tag == qn("w:p"):
            text = "".join(x.text or "" for x in prev.iter(qn("w:t"))).strip()
            has_break = (prev.find(qn("w:pPr")) is not None and
                         prev.find(qn("w:pPr")).find(qn("w:sectPr")) is not None)
            has_img = any(True for _ in prev.iter(qn("w:drawing")))
            # A field's begin/separate/end markers carry NO w:t, so a paragraph
            # holding one looks empty. Deleting those tore the List of Figures
            # field apart: its begin marker went, its cached result stayed, and
            # the page rendered the literal string TOC 	 "Figure Title,2".
            has_field = any(True for _ in prev.iter(qn("w:fldChar"))) or                 any(True for _ in prev.iter(qn("w:instrText")))
            if text or has_break or has_img or has_field:
                break
            nxt = prev.getprevious()
            prev.getparent().remove(prev)
            removed += 1
            prev = nxt
    return removed


def drop_empty_headings(doc) -> int:
    """Remove body-level heading paragraphs that contain nothing.

    The template leaves an empty Heading 2 where its List of Appendices used to
    start; followed by a page-break paragraph it produced a blank page x
    between the contents and the List of Tables. An empty heading is also an
    accessibility defect in its own right -- the checker reports it, and a
    screen reader announces a heading with no name.

    Paragraphs carrying a section break are never removed: one of them is an
    empty Heading 2, and it is what switches the body to arabic numbering.
    It is restyled as Body instead -- the break lives in its sectPr, not in
    its style -- because an empty heading is the defect either way."""
    from docx.oxml.ns import qn
    removed = 0
    for par in list(doc.element.body.iter(qn("w:p"))):
        if par.getparent() is not doc.element.body:
            continue
        pPr = par.find(qn("w:pPr"))
        if pPr is None:
            continue
        ps = pPr.find(qn("w:pStyle"))
        style = (ps.get(qn("w:val")) or "") if ps is not None else ""
        if not style.replace(" ", "").lower().startswith("heading"):
            continue
        if "".join(x.text or "" for x in par.iter(qn("w:t"))).strip():
            continue
        if any(True for _ in par.iter(qn("w:drawing"))) or \
                any(True for _ in par.iter(qn("w:fldChar"))):
            continue
        if pPr.find(qn("w:sectPr")) is not None:
            ps.set(qn("w:val"), doc.styles["Body"].style_id)
        else:
            par.getparent().remove(par)
        removed += 1
    return removed


def collapse_padding_runs(doc) -> int:
    """Collapse runs of three or more empty paragraphs down to one.

    Applied only AFTER the second section break, so the title page and the
    Certification of Approval keep the blank paragraphs they use for layout --
    the signature blocks in particular. What it removes is the template's
    padding after the Table of Contents, which produced a blank page x between
    the contents and the List of Tables."""
    from docx.oxml.ns import qn
    seen_breaks = 0
    run: list = []
    removed = 0
    for par in list(doc.element.body.iter(qn("w:p"))):
        # body.iter() descends into tables, and a w:tc MUST contain at least
        # one w:p. Emptying a table cell produced a file Word refused to open
        # with "The file appears to be corrupted", so only top-level
        # paragraphs are eligible.
        if par.getparent() is not doc.element.body:
            continue
        pPr = par.find(qn("w:pPr"))
        if pPr is not None and pPr.find(qn("w:sectPr")) is not None:
            seen_breaks += 1
            run = []
            continue
        if seen_breaks < 2:
            continue
        text = "".join(x.text or "" for x in par.iter(qn("w:t"))).strip()
        has_content = (text
                       or any(True for _ in par.iter(qn("w:drawing")))
                       or any(True for _ in par.iter(qn("w:fldChar")))
                       or any(True for _ in par.iter(qn("w:instrText")))
                       or any(True for _ in par.iter(qn("w:br"))))
        if has_content:
            run = []
            continue
        run.append(par)
        if len(run) > 1:
            par.getparent().remove(par)
            removed += 1
    return removed


def fix_front_matter_page_numbers(doc) -> str:
    r"""The template numbers its own certification and abstract pages in ARABIC.

    Section 0 (title, copyright) is lowerRoman, and section 2 (acknowledgements
    onward) restarts lowerRoman at 5 -- but section 1, which holds the
    Certification of Approval and the Abstract, is set to start at 1 with no
    format, which Word reads as arabic. Nobody noticed because the checklist
    says those pages carry no visible number: the template's own comments call
    them "Counted as Roman Numeral ii/iii, but not typed".

    They become visible the moment a Table of Contents references them, which
    is how "Certification of Approval ... 2" and "Abstract ... 3" appeared
    between "Acknowledgements ... v" and "Table of Contents ... vi".

    An earlier attempt fixed the symptom instead, setting w:outlineLvl to 9 to
    drop those entries from the contents. It had no effect and reported success:
    a TOC built with the \o switch selects by HEADING STYLE, not by outline
    level. Recorded because a fix that silently does nothing is worse than none.

    Title = i, copyright = ii, so section 1 starts at iii.
    """
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    sec = doc.sections[1]._sectPr
    pg = sec.find(qn("w:pgNumType"))
    if pg is None:
        pg = OxmlElement("w:pgNumType")
        sec.append(pg)
    pg.set(qn("w:fmt"), "lowerRoman")
    pg.set(qn("w:start"), "3")
    return "section 1 -> lowerRoman from iii"


def para_by_text(doc, needle):
    for p in doc.paragraphs:
        if needle.lower() in p.text.lower():
            return p
    return None


def set_text(par, text, bold=False):
    for r in list(par.runs):
        r._r.getparent().remove(r._r)
    r = par.add_run(text)
    r.bold = bold
    return par


def insert_after(par, text, style_obj):
    new = par.insert_paragraph_before(text)      # placed before, then moved
    par._p.addnext(new._p)
    new.style = style_obj
    for r in list(new.runs):
        r._r.getparent().remove(r._r)
    add_inline(new, text)
    return new


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default="FaultLines_thesis_SFSU.docx")
    a = ap.parse_args()

    try:
        import docx  # noqa: F401
        from docx import Document
    except ImportError:
        print("needs `pip install python-docx` in the project venv")
        return 1

    if not TEMPLATE.exists():
        print(f"fetching the official template -> {TEMPLATE.name}")
        try:
            with urllib.request.urlopen(TEMPLATE_URL, timeout=60) as r, \
                    open(TEMPLATE, "wb") as f:
                shutil.copyfileobj(r, f)
        except Exception as e:                       # noqa: BLE001
            print(f"could not download the SFSU template: {e}\n"
                  f"Download it by hand from\n  {TEMPLATE_URL}\n"
                  f"and save it as {TEMPLATE}")
            return 1

    doc = Document(str(TEMPLATE))
    paras = doc.paragraphs

    # ---------------- the template's placeholders and comments ---------
    # Fill the content controls where they stand. Nothing else on the title,
    # copyright or certification pages is touched -- the template's sentences
    # and spacing are already correct, and rewriting them is what broke the
    # first build.
    n_ctrl = fill_content_controls(doc)
    n_comments = strip_comments(doc)

    # ---------------- abstract + acknowledgements ----------------
    front_md = (WRITEUP / FRONT).read_text(encoding="utf-8")
    sections = re.split(r"^#\s+", front_md, flags=re.MULTILINE)
    named = {}
    for sec in sections:
        if not sec.strip():
            continue
        head, _, rest = sec.partition("\n")
        named[head.strip().lower()] = rest.strip()

    body_style = doc.styles["Body"]

    ab = para_by_text(doc, "Begin your abstract here")
    if ab is not None and "abstract" in named:
        parts = [b for b in blocks(named["abstract"]) if b[0] == "p"]
        set_text(ab, "")
        add_inline(ab, parts[0][1] if parts else "")
        ab.style = body_style
        anchor = ab
        for _k, text in parts[1:]:
            anchor = insert_after(anchor, text, body_style)

    ack = para_by_text(doc, "Normally the Acknowledgements section")
    if ack is not None and "acknowledgements" in named:
        parts = [b for b in blocks(named["acknowledgements"]) if b[0] == "p"]
        set_text(ack, "")
        add_inline(ack, parts[0][1] if parts else "")
        ack.style = body_style
        ack.alignment = 3                      # checklist: justified
        anchor = ack
        for _k, text in parts[1:]:
            anchor = insert_after(anchor, text, body_style)
            anchor.alignment = 3

    # ---------------- clear the lorem-ipsum body ----------------
    # Paragraph 114 carries the section break into the arabic-numbered body and
    # must survive; everything after it is placeholder content.
    start = None
    for idx, p in enumerate(doc.paragraphs):
        if p.text.strip() == "Introduction" and p.style.name == "Heading 2":
            start = idx
            break
    if start is None:
        print("could not find the template's 'Introduction' placeholder")
        return 1
    body_el = doc.element.body
    for p in list(doc.paragraphs)[start:]:
        p._p.getparent().remove(p._p)
    for tbl in list(doc.tables):
        # the placeholder tables live in the body section only
        if tbl._tbl.getparent() is body_el:
            tbl._tbl.getparent().remove(tbl._tbl)

    # ---------------- the chapters ----------------
    report: list = []
    em = Emitter(doc, report)
    missing = []
    for fname, title in CHAPTERS:
        path = WRITEUP / fname
        if not path.exists():
            missing.append(fname)
            continue
        md = path.read_text(encoding="utf-8")
        for kind, payload in blocks(md):
            em.emit(kind, payload)
        em.flush_pending_figure()
    if missing:
        print("missing chapters: " + ", ".join(missing))
        return 1

    removed, cleared = strip_template_furniture(doc)
    n_blank = collapse_blank_before_section_breaks(doc)
    n_pad = collapse_padding_runs(doc)
    n_eh = drop_empty_headings(doc)
    n_toc = fix_front_matter_page_numbers(doc)

    out = WRITEUP / a.out
    doc.save(str(out))

    if em.pending_ttitle:
        print(f"table title with no table after it: {em.pending_ttitle!r}")
        return 1

    # generated table titles are placeholders and say so, in a file
    stubs = WRITEUP / "TABLE_TITLES_TO_FIX.md"
    lines = ["# Table titles that still need a human",
             "",
             "`build_thesis_docx.py` generated one title per table from the "
             "nearest heading, because the SFSU checklist requires a numbered "
             "title above every table and inventing a good one is a judgement "
             "a script does not have.",
             "",
             "**These are placeholders.** Replace the part after the number "
             "with something that says what the table shows.",
             "",
             "| # | generated title |", "|---|---|"]
    for num, head in report:
        lines.append(f"| {num} | Table {num}. {head} |")
    if report:
        stubs.write_text("\n".join(lines) + "\n", encoding="utf-8")
    elif stubs.exists():
        stubs.unlink()

    print(f"-> {out.name}   {out.stat().st_size/1024/1024:.2f} MB")
    print(f"   {em.fig} figures (alt text set) - {em.tab} tables")
    if MISSING_FIGURES:
        print(f"   !! {len(MISSING_FIGURES)} FIGURE(S) MISSING - the "
              f"document has placeholders where images belong.")
        print("      run: python scripts/capture_figures.py")
    print(f"   template furniture: {removed} paragraphs removed, "
          f"{cleared} cleared (they carry section breaks)")
    print(f"   {n_ctrl} content controls filled - "
          f"{n_comments} comment anchors deleted - "
          f"front matter: {n_toc}")
    print(f"   {n_blank} padding paragraphs removed before section breaks, "
          f"{n_pad} in runs after the front matter, "
          f"{n_eh} empty headings")
    if report:
        print(f"   !! {len(report)} table(s) have no <!-- table: --> title "
              f"in the source - stubs listed in TABLE_TITLES_TO_FIX.md")
    else:
        print("   every table has its title from the source")
    print()
    print("   NEXT, EVERY TIME:")
    print("     python scripts/finalize_docx.py   (needs Word + pywin32) -")
    print("     updates the ToC / List of Tables / Figures and checks alt")
    print("     text, table titles, header rows and headings")
    print("   BEFORE SUBMITTING, BY HAND:")
    print("     - Review > Check Accessibility, as a final confirmation")
    print("     - the committee names in COMMITTEE, if still placeholders")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
