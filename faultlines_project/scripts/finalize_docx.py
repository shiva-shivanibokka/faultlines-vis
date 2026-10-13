#!/usr/bin/env python
"""Do in Word what build_thesis_docx.py cannot, then check the result.

    python scripts/finalize_docx.py [--pdf out.pdf]

Needs Microsoft Word and pywin32 (the anaconda base interpreter has it; the
project venv does not). Run it AFTER build_thesis_docx.py, every time: a
rebuild writes fresh, un-updated fields.

WHAT IT DOES
  1. Opens the .docx in an invisible Word, lays it out, and updates every
     field: the Table of Contents, the List of Tables, the List of Figures and
     the page numbers inside them. This is the Ctrl+A, F9 step. Page numbers
     do not exist until Word has paginated, which is why python-docx cannot.
  2. Saves, and optionally exports a PDF through Word so the layout can be
     looked at rather than assumed.
  3. Re-opens the saved file with python-docx and checks the things the SFSU
     checklist and Word's Accessibility Checker flag, so a failure is printed
     rather than found by the format reviewer:
       - every image has alt text
       - every table has a numbered Table Title directly above it and a
         marked header row
       - no empty headings, and no heading that skips a level
       - the List of Tables / Figures hold one entry per table / figure
       - the document has a title property
     Word's own Accessibility Checker is not exposed to automation, so this is
     a reimplementation of its error-level rules, not the checker itself.

Exits non-zero if any check fails.
"""
from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
DOCX = ROOT / "writeup" / "FaultLines_thesis_SFSU.docx"
TITLE = ("Detecting and Mapping Silent Consensus: "
         "Correlated Blind Spots in LLM Ensembles")

WD_FORMAT_PDF = 17


def update_in_word(path: Path, pdf: Path | None) -> dict:
    import win32com.client  # noqa: import here so --check-only needs no Word
    word = win32com.client.DispatchEx("Word.Application")
    word.Visible = False
    word.DisplayAlerts = 0
    try:
        doc = word.Documents.Open(str(path), False, False, False)
        doc.Repaginate()
        # Fields first so the captions and page numbers exist, then the lists
        # built from them, then fields again for any PAGEREF inside the lists.
        doc.Fields.Update()
        for toc in doc.TablesOfContents:
            toc.Update()
        for tof in doc.TablesOfFigures:
            tof.Update()
        doc.Repaginate()
        doc.Fields.Update()
        doc.BuiltInDocumentProperties("Title").Value = TITLE
        stats = {
            "pages": doc.ComputeStatistics(2),       # wdStatisticPages
            "tocs": doc.TablesOfContents.Count,
            "tofs": doc.TablesOfFigures.Count,
        }
        doc.Save()
        if pdf:
            doc.ExportAsFixedFormat(str(pdf), WD_FORMAT_PDF)
        doc.Close(False)
        return stats
    finally:
        word.Quit()


def check(path: Path, n_tables_expected: int | None = None) -> list[str]:
    import docx
    from docx.oxml.ns import qn
    d = docx.Document(str(path))
    problems: list[str] = []

    # images: alt text lives on wp:docPr/@descr
    body = d.element.body
    for i, pr in enumerate(body.iter(qn("wp:docPr")), 1):
        if not (pr.get("descr") or "").strip():
            problems.append(f"image {i} ({pr.get('name')}) has no alt text")

    # walk body children in order, so "the paragraph before a table" is real
    prev = None
    n_tab = 0
    for el in body.iterchildren():
        if el.tag == qn("w:tbl"):
            # template furniture tables (title page etc.) have no title and
            # sit before the body; only tables after the first Table Title
            # paragraph are thesis tables
            ptxt, pstyle = prev if prev else ("", "")
            if pstyle == "Table Title":
                n_tab += 1
                if not re.match(rf"Table {n_tab}\. \S", ptxt):
                    problems.append(f"table {n_tab}: title reads {ptxt!r}")
                first = el.find(qn("w:tr"))
                trpr = first.find(qn("w:trPr")) if first is not None else None
                if trpr is None or trpr.find(qn("w:tblHeader")) is None:
                    problems.append(f"table {n_tab}: header row not marked")
            elif n_tab:
                problems.append(f"a table after Table {n_tab} has no title")
            prev = None
        elif el.tag == qn("w:p"):
            p = docx.text.paragraph.Paragraph(el, d)
            txt = p.text.strip()
            if txt or el.find(".//" + qn("w:drawing")) is not None:
                prev = (txt, p.style.name if p.style is not None else "")
    if n_tables_expected is not None and n_tab != n_tables_expected:
        problems.append(f"{n_tab} titled tables, expected {n_tables_expected}")

    # headings: none empty, none skipping a level (after the front matter)
    last = None
    for p in d.paragraphs:
        name = p.style.name if p.style is not None else ""
        m = re.fullmatch(r"Heading (\d)", name)
        if not m:
            continue
        lvl = int(m.group(1))
        if not p.text.strip():
            problems.append(f"empty {name}")
            continue
        if last is not None and lvl > last + 1:
            problems.append(f"heading skips a level: {name} after Heading "
                            f"{last}: {p.text.strip()[:60]!r}")
        last = lvl

    # the lists: count entries Word wrote into the ToF fields
    lot = sum(1 for p in d.paragraphs
              if re.match(r"Table \d+\. .*\t\d+$", p.text))
    lof = sum(1 for p in d.paragraphs
              if re.match(r"Figure \d+\. .*\t\d+$", p.text))
    toc = sum(1 for p in d.paragraphs
              if p.style is not None
              and p.style.name.lower().startswith("toc")
              and re.search(r"\t[ivx\d]+$", p.text))
    n_fig = sum(1 for p in d.paragraphs
                if p.style is not None and p.style.name == "Figure Title")
    if lot != n_tab:
        problems.append(f"List of Tables has {lot} entries, {n_tab} tables")
    if lof != n_fig:
        problems.append(f"List of Figures has {lof} entries, {n_fig} figures")
    if toc == 0:
        problems.append("Table of Contents has no entries with page numbers")

    if not (d.core_properties.title or "").strip():
        problems.append("document has no title property")

    print(f"   {n_tab} tables, {n_fig} figures, "
          f"ToC {toc} entries, LoT {lot}, LoF {lof}")
    return problems


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--docx", type=Path, default=DOCX)
    ap.add_argument("--pdf", type=Path, default=None,
                    help="also export a PDF through Word")
    ap.add_argument("--check-only", action="store_true",
                    help="skip Word; only run the checks")
    a = ap.parse_args()
    path = a.docx.resolve()
    if not a.check_only:
        stats = update_in_word(path, a.pdf.resolve() if a.pdf else None)
        print(f"-> fields updated in Word: {stats['pages']} pages, "
              f"{stats['tocs']} ToC, {stats['tofs']} lists of tables/figures")
        if a.pdf:
            print(f"-> {a.pdf}")
    problems = check(path)
    for p in problems:
        print(f"   !! {p}")
    print("   all checks pass" if not problems
          else f"   {len(problems)} problem(s)")
    return 1 if problems else 0


if __name__ == "__main__":
    sys.exit(main())
